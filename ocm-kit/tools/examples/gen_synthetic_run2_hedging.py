#!/usr/bin/env python3
# ============================================================================
# EXAMPLE ONLY -- THIS IS RUN 2's GENERATOR, NOT A KIT TOOL.
#
# It produces run 2's hedging inputs: EUR swap curves, a 1899-12-30 spreadsheet
# epoch, maturities 1..30/35/40/45/50/60, observation dates 2012-01..2026-03.
# Running it for any other migration yields inputs the source cannot read.
#
# Synthetic generation is irreducibly source-specific, so the kit cannot ship a
# general one. Each run writes its own and keeps it beside its evidence. This
# file is retained under tools/examples/ purely as a worked shape: deterministic
# seed, asserted invariants, every layout fact recorded as an unverified
# assumption against the real files.
#
# Whatever a stand-in shows you is tagged `observed_in_synthetic_input`, never
# `observed_in_real_input` (lessons, finding 3), and caps the verdict under R2.
# ============================================================================
# Controller tool: generate synthetic in-format inputs (provenance: synthetic, R2).
# Deterministic. Formats mirror the source system's documented input shapes;
# every layout fact here is recorded as an unverified assumption against the
# real files (see input_profile.json).
import random, math, datetime, os, sys
import sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "inputs"
random.seed(20260828)
D0 = datetime.date(1899, 12, 30)          # spreadsheet epoch

# ---------- observation dates: last weekday of month, 2012-01 .. 2026-03 ----
def last_weekday(y, m):
    d = datetime.date(y + (m == 12), (m % 12) + 1, 1) - datetime.timedelta(days=1)
    while d.weekday() > 4:
        d -= datetime.timedelta(days=1)
    return d

dates = []
y, m = 2012, 1
while (y, m) <= (2026, 3):
    dates.append(last_weekday(y, m))
    m += 1
    if m == 13: y, m = y + 1, 1
assert len(dates) == 171, len(dates)
must_have = ["2012-01-31","2014-01-31","2016-01-29","2018-01-31","2020-01-31",
             "2024-01-31","2019-12-31","2022-03-31","2026-03-31","2019-01-31"]
have = {d.isoformat() for d in dates}
missing = [x for x in must_have if x not in have]
assert not missing, missing

# ---------- swap curve model (plausible EUR path, R4) -----------------------
MATS = list(range(1, 31)) + [35, 40, 45, 50, 60]      # gapped set (R3, in full)

def interp(anchors, t):
    ks = sorted(anchors)
    if t <= ks[0]: return anchors[ks[0]]
    for a, b in zip(ks, ks[1:]):
        if t <= b:
            w = (t - a) / (b - a)
            return anchors[a] * (1 - w) + anchors[b] * w
    return anchors[ks[-1]]

def tof(d): return d.year + (d.month - 1) / 12.0

SHORT = {2012.0: 1.20, 2014.0: 0.40, 2016.0: -0.10, 2019.0: -0.35, 2020.6: -0.52,
         2021.9: -0.45, 2022.4: 0.60, 2023.4: 3.05, 2024.4: 3.30, 2025.4: 2.40, 2026.2: 2.10}
SPREAD = {2012.0: 1.60, 2016.0: 1.30, 2020.0: 0.85, 2022.0: 0.95, 2023.5: -0.25,
          2025.0: 0.30, 2026.2: 0.50}
HUMP = {2012.0: 0.10, 2018.0: 0.06, 2022.0: -0.05, 2024.0: 0.12, 2026.2: 0.08}

def rate(d, mat):
    t = tof(d)
    s, sp, h = interp(SHORT, t), interp(SPREAD, t), interp(HUMP, t)
    r = s + sp * (1 - math.exp(-mat / 9.0)) + h * (mat / 6.0) * math.exp(1 - mat / 6.0)
    return round(r, 3)

def de(v, nd=3):
    return ("%.*f" % (nd, v)).replace(".", ",")

# ---------- loans -----------------------------------------------------------
YEARS = list(range(2021, 2045))                      # 24 maturity years exactly
wt = {yy: 60 + max(0, 10 - abs(yy - 2030)) * 28 for yy in YEARS}
wt[2030] += 260
N_LOANS = 9000
year_pool = [yy for yy in YEARS for _ in range(wt[yy])]

loans = []
bsv_used = set()
def new_bsv():
    while True:
        b = random.randint(1000000, 9999999)
        if b not in bsv_used:
            bsv_used.add(b); return b

for i in range(N_LOANS):
    yy = random.choice(year_pool)
    mm, dd = random.randint(1, 12), random.choice([5, 15, 28])
    zb = yy * 10000 + mm * 100 + dd
    beg_y = yy - random.choice([8, 9, 10, 10, 11, 12])
    vb = beg_y * 10000 + random.randint(1, 12) * 100 + random.choice([5, 15, 28])
    da = round(random.uniform(5000, 400000), 2)
    zins = round(random.uniform(1.2, 5.8), 2)
    unter = random.choices([1, 2, 3], weights=[75, 15, 10])[0]
    ratio = random.gauss(0.95, 0.15)
    if random.random() < 0.02: ratio = random.uniform(2.2, 8.0)   # data problems
    ratio = min(max(ratio, 0.45), 8.0)
    teuro = round(da / 1000.0 / ratio, 1)
    if teuro <= 0: teuro = 0.1
    loans.append(dict(BSV=new_bsv(), DaBetrag=da, Zins=zins, Zinsbindung=zb,
                      Vertragsbeginn=vb, unterlegt=unter, teuro=teuro))

for l in random.sample(loans, 15):                    # a few missing rates
    l["Zins"] = None
dups = random.sample(loans, 40)                       # duplicated BSV rows
for l in dups:
    c = dict(l); c["DaBetrag"] = round(l["DaBetrag"] * random.uniform(0.4, 0.9), 2)
    loans.append(c)
orphans = set(l["BSV"] for l in random.sample(loans, 200))   # not in KOBRA -> merge drops
random.shuffle(loans)

# ---------- xlsx writing ----------------------------------------------------
from openpyxl import Workbook

def write_market(path):
    wb = Workbook(); ws = wb.active; ws.title = "Historical Swap Data"
    ws["B1"] = "Start Date"; ws["C1"] = (dates[0] - D0).days
    ws["B2"] = "End Date";   ws["C2"] = (dates[-1] - D0).days
    ws["B4"] = "Quelle";     ws["C4"] = "BLOOMBERG"
    ws["B6"] = "Dates"
    for j, d in enumerate(dates):
        ws.cell(row=6, column=5 + j, value=(d - D0).days)
    instruments = ["EONIA Index"] + ["EUSA%d Curncy" % mmm for mmm in MATS]
    for i, ins in enumerate(instruments):
        r = 7 + i
        ws.cell(row=r, column=2, value=ins)
        ws.cell(row=r, column=3, value="#NAME?")
        ws.cell(row=r, column=4, value="PX_LAST")
        mat = 0.02 if i == 0 else MATS[i - 1]
        for j, d in enumerate(dates):
            if ins == "EUSA25 Curncy" and d.isoformat() in ("2013-03-29", "2013-04-30", "2013-05-31"):
                ws.cell(row=r, column=5 + j, value="#NAME?"); continue
            if ins == "EUSA18 Curncy" and d.isoformat() == "2013-06-28":
                continue                                # blank cell -> NA
            v = (interp(SHORT, tof(d)) - 0.05) if i == 0 else rate(d, mat)
            ws.cell(row=r, column=5 + j, value=de(v))
    info = wb.create_sheet("Info"); info["A1"] = "Export"; info["B1"] = "MarktdatenExport"
    wb.save(path)

def write_loans(path):
    wb = Workbook(); ws = wb.active; ws.title = "Daten"
    ws.append(["BSV", "DaBetrag", "Zins", "Zinsbindung", "Vertragsbeginn",
               "unterlegt als", "Tarif", "Notiz"])
    for l in loans:
        ws.append([l["BSV"], l["DaBetrag"], l["Zins"], l["Zinsbindung"],
                   l["Vertragsbeginn"], l["unterlegt"],
                   "T%d" % random.randint(1, 6), "ok"])
    wb.save(path)

# ---------- KOBRA text export ----------------------------------------------
def de_th(v, nd=2):
    s = "%.*f" % (nd, v); ip, fp = s.split(".")
    sign = ""
    if ip.startswith("-"): sign, ip = "-", ip[1:]
    grp = ""
    while len(ip) > 3:
        grp = "." + ip[-3:] + grp; ip = ip[:-3]
    return sign + ip + grp + "," + fp

def ymd(y_, m_, d_): return "%04d%02d%02d" % (y_, m_, d_)

def kobra_row(vnr, teuro, has_loan_flag, rng):
    ab_y = rng.randint(1995, 2023)
    ab = ymd(ab_y, rng.randint(1, 12), rng.choice([2, 11, 21]))
    st = ymd(ab_y, rng.randint(1, 12), 1)
    ein = ymd(min(ab_y + 1, 2024), rng.randint(1, 12), 15) if rng.random() < 0.9 else "0"
    terminated = rng.random() < 0.08
    kdat = ymd(rng.randint(max(ab_y + 1, 2005), 2024), rng.randint(1, 12), 28) if terminated else "0"
    if terminated:
        status = rng.choice([290, 324, 280])
    elif has_loan_flag:
        status = 24
    else:
        status = rng.choice([1, 18, 18, 52, 84, 100])
    guth = rng.uniform(0, 60000) if not has_loan_flag else rng.uniform(0, 900)
    best = rng.uniform(3000, 350000) if has_loan_flag else 0.0
    f = [""] * 80
    f[0], f[1], f[2] = ab, st, ein
    f[3] = str(status)
    f[4] = de(teuro, 1)                               # simple decimal comma
    f[5] = de_th(guth)                                # thousands separator -> text column
    f[6] = de(rng.uniform(0, 300), 2); f[7] = de(rng.uniform(0, 80), 2)
    f[8] = de(rng.uniform(0, 40), 2);  f[9] = de(rng.uniform(0, 20), 2)
    f[10] = de(rng.uniform(0, 900), 1); f[11] = de_th(rng.uniform(0, 90000))
    f[12] = ymd(rng.randint(2010, 2024), rng.randint(1, 12), 1) if rng.random() < 0.5 else "0"
    f[13] = de_th(rng.uniform(0, 250000)) if has_loan_flag else "0,00"
    f[14] = de_th(rng.uniform(0, 50000))
    f[15] = de_th(best)
    f[16] = de(rng.uniform(0, 4000), 2); f[17] = de(rng.uniform(0, 800), 2)
    f[18] = de(rng.uniform(0, 120), 2);  f[19] = de(rng.uniform(0, 5000), 2)
    f[20] = de_th(best * rng.uniform(0.9, 1.1)) if has_loan_flag else "0,00"
    f[21] = de(rng.uniform(1.0, 4.5), 2); f[22] = de(rng.uniform(1.0, 4.5), 2)
    f[23] = "T%d" % rng.randint(1, 6)
    f[24] = str(vnr)
    f[30] = rng.choice(["I0", "I1", ""])
    f[31] = rng.choice(["O1", ""]); f[32] = rng.choice(["J", "N", ""])
    f[33] = rng.choice(["Z1", "Z2", ""]); f[34] = de(rng.uniform(0, 900), 2)
    for k in range(35, 59):
        f[k] = de(rng.uniform(0, 500), 2) if rng.random() < 0.8 else "0,00"
    f[59] = rng.choice(["L", ""])
    f[60] = de(rng.uniform(0, 700), 2); f[61] = de(rng.uniform(0, 40), 2)
    f[62] = de(rng.uniform(0, 470), 2)
    f[64] = kdat
    f[65] = de_th(rng.uniform(400, 20000)); f[66] = de(rng.uniform(0, 100), 1)
    f[67] = rng.choice(["", "Z"]); f[68] = rng.choice(["M1", "M2"])
    f[69] = ymd(rng.randint(2000, 2030), rng.randint(1, 12), 1)
    f[70] = rng.choice(["KO1", "KO2"]); f[71] = de(rng.uniform(0, 10), 1)
    f[72] = "ZB %s" % de(rng.uniform(1, 5), 2)
    f[73] = rng.choice(["A", "B", ""])
    f[74] = ymd(rng.randint(2010, 2024), rng.randint(1, 12), 1) if rng.random() < 0.4 else "0"
    f[75] = rng.choice(["W1", "W2", ""])
    f[76] = rng.choice(["A", ""])
    f[77] = "0"
    f[78] = rng.choice(["W0", ""])
    f[79] = rng.choice(["G1", "G2", "G3", "G4"])
    return f

def write_kobra(path, n_total, seed, loan_link):
    rng = random.Random(seed)
    rows = []
    if loan_link:
        for l in loans:
            if l["BSV"] in orphans:  continue
            rows.append(kobra_row(l["BSV"], l["teuro"], True, rng))
    seen = set(int(r[24]) for r in rows)
    while len(rows) < n_total:
        v = rng.randint(1000000, 99999999)
        if v in seen: continue
        seen.add(v)
        rows.append(kobra_row(v, rng.uniform(5, 400), rng.random() < 0.25, rng))
    for r in rng.sample(rows, 12):                    # duplicate contract numbers
        rows.append(list(r))
    ragged = rng.sample(range(len(rows)), 30)         # short lines (fill semantics)
    rng.shuffle(rows)
    with open(path, "w", encoding="iso-8859-1", newline="\n") as fh:
        for i, r in enumerate(rows):
            rr = r[:78] if i in ragged else r
            fh.write("|".join(rr) + "\n")
    return len(rows)

os.makedirs(OUT, exist_ok=True)
write_market(os.path.join(OUT, "MarktdatenExport copy.xlsx"))
write_loans(os.path.join(OUT, "Kopie von Darlehen zum BSV.xlsx"))
n1 = write_kobra(os.path.join(OUT, "kobra202412.txt"), 60000, 7, True)
n2 = write_kobra(os.path.join(OUT, "kobra202412_full.txt"), 865524, 11, True)
print("dates:", len(dates), "| loans rows:", len(loans),
      "| kobra rows:", n1, "| kobra full rows:", n2)
print("maturities:", MATS)
